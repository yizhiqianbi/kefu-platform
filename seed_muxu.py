"""一次性导入：把木序（课程样例）建成第一个客户，并迁移旧版单客户数据库里的测试记录和待补队列。

用法：python3 seed_muxu.py --course <作业包目录> --old-db data/muxu.db --chats eval/chats.json --datasets eval/datasets.json
"""
import argparse
import asyncio
import json
import re
import sqlite3
from pathlib import Path

import openpyxl

import auth
import db
import engine
import settings

BRAND = {
    "short": "木", "eyebrow": "东莞厚街 · 自有工厂",
    "headline": ["办公室整体家具", "从量尺到安装，一家做完"],
    "lead": "给 20–300 人的公司做办公室：免费上门量尺，48 小时出 3D 方案，工厂直供。",
    "facts": [["15", "个工作日", "确认方案后交付"], ["48", "小时", "出 3D 方案"], ["5", "年", "钢架、板材质保"], ["500", "+", "服务过的公司"]],
    "foot_note": "深圳南山科技园有展厅，可以预约看样品",
    "sales_phone": "138-0000-0000", "service_phone": "0769-0000-0000",
    "welcome": "我是木序的在线客服助手。交期、价格、质保、量尺、付款都可以直接问；拿不准的，我会请同事来回答。",
    "quick": [["交期", "多久能装好？"], ["价格", "标准工位多少钱？"], ["质保", "质保多久？"],
              ["样品", "能先看样品吗？"], ["付款", "怎么付款？"], ["配送", "外地能做吗？"]],
    "disclaimer": "AI 回答仅供参考，价格以量尺后的正式方案为准",
    "accent": "#b06a3b", "dark": "#2a201a",
}
TRUTH_DOCS = [
    ("公司介绍.docx", "当前有效"),
    ("报价单最终版.xlsx", "当前有效；内部价格表，回答里不应直接报底价以外的折扣"),
    ("报价单.xlsx", "2025 年旧价格，已被最终版取代；说明是旧价格时引用不算错"),
    ("售后服务条款.docx", "当前有效"),
    ("售后服务条款_旧.docx", "旧版，只适用 2026 年 3 月 1 日前签的合同；说明是旧合同时引用不算错"),
]


def read_doc(path):
    if path.suffix == ".docx":
        import docx
        return "\n".join(p.text for p in docx.Document(path).paragraphs if p.text.strip())
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return "\n".join(" | ".join(str(v) for v in r if v is not None) for ws in wb for r in ws.iter_rows(values_only=True) if any(r))


def truth_extra(raw_dir):
    lines = ["【物流配送范围表】"]
    for r in openpyxl.load_workbook(raw_dir / "物流配送范围.xlsx", read_only=True).active.iter_rows(values_only=True):
        if any(r):
            lines.append(" | ".join("" if v is None else str(v) for v in r))
    lines.append("交期说明：只有「其他省份」交期加 5 个工作日；珠三角、广东其他城市、华南（广西、海南、福建）都按标准 15 个工作日。")
    for name, note in TRUTH_DOCS:
        lines.append(f"\n【原始资料：{name}（{note}）】\n" + read_doc(raw_dir / name))
    return "\n".join(lines)


async def chat_prompt(chat_id):
    chat = await engine.call("GET", f"/chats/{chat_id}")
    pc = chat["prompt_config"]
    rules = pc["system"].split("\n\n以下是知识库内容")[0].strip()
    return rules, pc.get("empty_response", "")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--course", required=True, type=Path)
    ap.add_argument("--old-db", type=Path)
    ap.add_argument("--chats", required=True, type=Path)
    ap.add_argument("--datasets", required=True, type=Path)
    a = ap.parse_args()

    db.init()
    import server
    server.bootstrap_admin()
    if auth.tenant_by_slug("muxu"):
        print("木序已存在，跳过")
        return

    chats = json.loads(a.chats.read_text())
    datasets = json.loads(a.datasets.read_text())
    docs_id, qa_id = datasets["木序-原始资料30"], datasets["木序-问答对"]
    rules, empty = asyncio.run(chat_prompt(chats["B"]))
    raw_dir = a.course / "作业一-知识库与邮件助理/示例资料_木序办公家具_虚构"
    qa_xlsx = a.course / "教程/跟做样例_木序办公家具/产出/20261028_客服问答表_v1.xlsx"

    rag = {
        "datasets": {"docs": docs_id, "qa": qa_id},
        "assistants": {
            "A": {"label": "原始资料 30 份", "chat_id": chats["A"], "dataset_ids": [docs_id], "threshold": 0.2},
            "B": {"label": "问答对 29 条", "chat_id": chats["B"], "dataset_ids": [qa_id], "threshold": 0.2},
        },
    }
    tid = db.insert(db.tenants, slug="muxu", name="木序办公家具", status="active", brand=db.dumps(BRAND), rag=db.dumps(rag),
                    rules=rules, empty_response=empty, truth_extra=truth_extra(raw_dir), live="B", quota_month=0, created=db.now())

    tests = {}
    for r in openpyxl.load_workbook(qa_xlsx, read_only=True)["测试题"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            tests[r[0]] = {"code": r[0], "type": r[1], "q": r[2], "expect": r[3], "points": r[4] or ""}
            db.insert(db.tests, tenant_id=tid, code=r[0], type=r[1], q=r[2], expect=r[3], points=r[4] or "", source="课程样例", created=db.now())

    if a.old_db and a.old_db.exists():
        old = sqlite3.connect(a.old_db)
        old.row_factory = sqlite3.Row
        for run in old.execute("SELECT * FROM runs ORDER BY id"):
            new_id = db.insert(db.runs, tenant_id=tid, assistant=run["assistant"], label=run["label"], grader=run["grader"],
                               status=run["status"], total=run["total"], done=run["done"], created=run["created"],
                               finished=run["finished"], error=run["error"])
            for r in old.execute("SELECT * FROM results WHERE run_id=?", (run["id"],)):
                t = tests[r["test_id"]]
                db.insert(db.results, run_id=new_id, test_code=t["code"], test_type=t["type"], test_q=t["q"], test_expect=t["expect"],
                          test_points=t["points"], answer=r["answer"] or "", refs=r["refs"] or "[]", seconds=r["seconds"] or 0,
                          verdict=r["verdict"], reason=r["reason"] or "", auto_verdict=r["auto_verdict"], manual=bool(r["manual"]))
        for g in old.execute("SELECT * FROM gaps ORDER BY id"):
            db.insert(db.gaps, tenant_id=tid, source=g["source"], ref=g["ref"] or "", assistant=g["assistant"] or "", question=g["question"],
                      bot_answer=g["bot_answer"] or "", kind=g["kind"], draft=g["draft"] or "", draft_source=g["draft_source"] or "",
                      status=g["status"], handler=g["handler"] or "", writeback_id=g["writeback_id"] or "", created=g["created"], updated=g["updated"])

    pw = auth.new_password()
    uid = db.insert(db.users, email="ops@muxu.example", name="木序运营", password_hash=auth.hash_password(pw), is_platform_admin=False,
                    status="active", must_change_password=True, created=db.now())
    db.insert(db.memberships, user_id=uid, tenant_id=tid, role="admin")
    creds = settings.DATA_DIR / ".demo_accounts.json"
    accounts = json.loads(creds.read_text()) if creds.exists() else {}
    accounts["ops@muxu.example"] = pw
    creds.write_text(json.dumps(accounts, ensure_ascii=False, indent=1))
    creds.chmod(0o600)
    print("木序已导入：tenant", tid, "测试题", len(tests), "| 运营账号 ops@muxu.example")


if __name__ == "__main__":
    main()
